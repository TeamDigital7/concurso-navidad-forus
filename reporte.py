"""Excel del concurso: hoja Dashboard con los indicadores, hoja Acumulado y una hoja por casa.

Se arma solo cuando alguien toca "Descargar" (el botón recibe la función, no el archivo).
No incluye premios: es información que no se muestra por tienda.
"""
import io
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

from concurso import BONUS_FIDELIZACION, CASAS, DESAFIOS_SSS, RANGO_PAGO
from puntos import factor_pago, resumen_casas

ASSETS = Path(__file__).parent / "assets"
PAGA_DESDE = RANGO_PAGO[0][0]

# Paleta: noche del castillo, burdeos de Gryffindor y oro del cáliz
NOCHE, PANEL, PANEL_2 = "0E1220", "182033", "212B42"
BURDEOS, ORO, ORO_CLARO = "5A1A1A", "C9A24D", "F1D791"
TEXTO, GRIS = "ECE6D6", "9AA3B5"
VERDE, ROJO = "3FBF7F", "E5534B"
PERGAMINO = "FAF5E8"

SERIF = "Georgia"  # serif de Windows y Mac: el toque de pergamino de los títulos


def _fill(c: str) -> PatternFill:
    return PatternFill("solid", start_color=c, end_color=c)


def _hex(casa: str) -> str:
    # Curiosity es casi negro: igual que en la app, se aclara para que se lea
    return "3B5498" if casa == "Curiosity" else CASAS[casa]["color"].lstrip("#")


# Tono claro de cada casa para texto sobre el fondo oscuro del Dashboard
CLARO = {"Speed": "5AA2FF", "Curiosity": "8EA2D8", "Ownership": "F3D778", "Accountability": "6FCB86",
         "Agility": "E58A8A"}


def _imagen(ruta: Path, ancho: int, formato: str = "PNG", lado: int | None = None) -> XLImage:
    """Imagen para Excel (no acepta webp): se convierte y se escala al ancho en píxeles."""
    im = Image.open(ruta).convert("RGBA" if formato == "PNG" else "RGB")
    if lado:
        im = im.resize((lado, lado), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, formato, quality=88)
    buf.seek(0)
    xl = XLImage(buf)
    xl.width, xl.height = ancho, round(ancho * im.height / im.width)
    return xl


def _escudo(casa: str, px: int) -> XLImage:
    nombre = CASAS[casa]["escudo"].rsplit(".", 1)[0]
    return _imagen(ASSETS.parent / "static" / "escudos" / "mini" / f"{nombre}-256.webp", px, lado=128)


def _celdas(ws, rango: str):
    """Las celdas de un rango ("B2:C3") o de una sola ("B2")."""
    if ":" not in rango:
        return [ws[rango]]
    return [c for fila in ws[rango] for c in fila]


def _pinta(ws, rango: str, color: str):
    for c in _celdas(ws, rango):
        c.fill = _fill(color)


def _celda(ws, ref: str, valor=None, *, size=11, bold=False, color=TEXTO, fill=None, fmt=None,
           h="left", v="center", font=None, merge: str | None = None, wrap=False):
    if merge:
        ws.merge_cells(merge)
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.font = Font(name=font or "Calibri", size=size, bold=bold, color=color)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if fill:
        for x in _celdas(ws, merge or ref):
            x.fill = _fill(fill)
    if fmt:
        c.number_format = fmt
    return c


def _num(x) -> float | None:
    """Excel no acepta NaN ni infinito: quedan vacíos."""
    if x is None or pd.isna(x) or (isinstance(x, float) and np.isinf(x)):
        return None
    return float(x)


def _texto_ejes(color: str = TEXTO) -> RichText:
    cp = CharacterProperties(sz=900, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _grafico_casas(ws, titulo: str, ref_cats, ref_vals, casas: list[str], fmt: str) -> BarChart:
    ch = BarChart()
    ch.type, ch.style, ch.legend = "col", 10, None
    ch.title = titulo
    ch.add_data(ref_vals, titles_from_data=False)
    ch.set_categories(ref_cats)
    s = ch.series[0]
    for i, casa in enumerate(casas):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = _hex(casa)
        dp.graphicalProperties.line.solidFill = ORO
        s.dPt.append(dp)
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    s.dLbls.numFmt = fmt
    s.dLbls.txPr = _texto_ejes(ORO_CLARO)
    ch.y_axis.numFmt = fmt
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="2A3550"))
    ch.y_axis.txPr = ch.x_axis.txPr = _texto_ejes(GRIS)
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(solidFill=ORO))
    ch.plot_area.graphicalProperties = GraphicalProperties(solidFill=PANEL)
    cp = CharacterProperties(sz=1200, b=True, solidFill=ORO_CLARO, latin=None)
    ch.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=cp)
    for r in ch.title.tx.rich.p[0].r:
        r.rPr = cp
    ch.width, ch.height = 15.5, 7.6
    return ch


# ---------- Hoja Dashboard ----------
def _dashboard(ws, rank: pd.DataFrame, dias: pd.DataFrame, competencias: list[dict], titulo: str):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ORO
    ws.column_dimensions["A"].width = 3
    for col in range(2, 14):
        ws.column_dimensions[get_column_letter(col)].width = 13.5
    ws.column_dimensions["N"].width = 3
    _pinta(ws, "A1:N78", NOCHE)

    # Banda: castillo, título y los escudos de las 5 casas (assets/excel-banda.jpg)
    banda = _imagen(ASSETS / "excel-banda.jpg", 1148, "JPEG")
    ws.add_image(banda, "B2")
    for f in range(1, 13):
        ws.row_dimensions[f].height = 15.75
    _celda(ws, "B13", titulo, size=10, color=GRIS, merge="B13:M13")

    # Indicadores del Canal Retail
    venta, ppto, ly = dias["venta"].sum(), dias["presupuesto"].sum(), dias["venta_ly"].sum()
    cumpl = venta / ppto if ppto else np.nan
    crec = dias["venta"].where(dias["venta_ly"] > 0, 0).sum() / ly - 1 if ly else np.nan
    fid = dias["tickets_fidelizados"].sum() / dias["tickets"].sum() if dias["tickets"].sum() else np.nan
    mes = dias["fecha"].max().month if len(dias) else 11
    meta_sss = DESAFIOS_SSS.get(mes, 0)
    f = factor_pago(cumpl) if ppto else 0
    paga_ok = (cumpl >= PAGA_DESDE) if ppto else False
    _celda(ws, "B15", "✦ CANAL RETAIL", size=15, bold=True, color=ORO_CLARO, font=SERIF, merge="B15:M15")
    tarjetas = [
        ("CUMPLIMIENTO", cumpl, "0.0%", VERDE if paga_ok else ROJO, f"se paga desde {PAGA_DESDE:.0%}"),
        ("SE PAGA", f, "0%", VERDE if f >= 1 else (ORO_CLARO if f else ROJO), "del premio"),
        ("VENTA", venta, '"S/ "#,##0.0,,"M"', TEXTO, f"presupuesto S/ {ppto / 1e6:,.1f}M"),
        ("SSS", crec, "+0.0%;-0.0%", VERDE if crec >= meta_sss else (ORO_CLARO if crec >= 0 else ROJO),
         f"meta +{meta_sss:.0%} · días comparables"),
        ("FIDELIZACIÓN", fid, "0.0%", VERDE if fid > BONUS_FIDELIZACION["umbral"] else ROJO,
         f"bonus sobre {BONUS_FIDELIZACION['umbral']:.0%}"),
        ("MAGOS ≥ 100%", int((rank["cumplimiento"] >= 1).sum()), "0", ORO_CLARO, f"de {len(rank)} tiendas"),
    ]
    borde = Side(style="medium", color=ORO)
    for i, (lab, val, fmt, color, sub) in enumerate(tarjetas):
        c1, c2 = get_column_letter(2 + 2 * i), get_column_letter(3 + 2 * i)
        _celda(ws, f"{c1}17", lab, size=9, bold=True, color=ORO, fill=PANEL, merge=f"{c1}17:{c2}17", h="center")
        _celda(ws, f"{c1}18", _num(val), size=22, bold=True, color=color, fill=PANEL, fmt=fmt,
               merge=f"{c1}18:{c2}19", h="center", font=SERIF)
        _celda(ws, f"{c1}20", sub, size=8, color=GRIS, fill=PANEL, merge=f"{c1}20:{c2}20", h="center")
        for f_ in range(17, 21):
            for col in (c1, c2):
                c = ws[f"{col}{f_}"]
                c.border = Border(top=borde if f_ == 17 else None, bottom=borde if f_ == 20 else None,
                                  left=borde if col == c1 else None, right=borde if col == c2 else None)
    ws.row_dimensions[18].height = ws.row_dimensions[19].height = 20

    # Las casas
    _celda(ws, "B23", "✦ LAS CASAS", size=15, bold=True, color=ORO_CLARO, font=SERIF, merge="B23:M23")
    casas = resumen_casas(rank).sort_values("cumplimiento", ascending=False)
    cab = ["CASA", "", "MAGOS", "VENTA S/", "", "PRESUPUESTO S/", "", "% PPTO", "SSS", "FIDELIZ.", "MAGOS ≥100%", ""]
    for i, t in enumerate(cab):
        _celda(ws, f"{get_column_letter(2 + i)}25", t, size=9, bold=True, color=ORO, fill=BURDEOS, h="center")
    for col in ("B25:C25", "E25:F25", "G25:H25", "L25:M25"):
        ws.merge_cells(col)
    for k, (_, r) in enumerate(casas.iterrows()):
        fila = 26 + k
        de_casa = rank[rank["casa"] == r["casa"]]
        fz = de_casa["tickets_fidelizados"].sum() / de_casa["tickets"].sum() if de_casa["tickets"].sum() else None
        fondo = PANEL if k % 2 == 0 else PANEL_2
        _celda(ws, f"B{fila}", r["casa"], bold=True, color="FFFFFF", fill=_hex(r["casa"]), merge=f"B{fila}:C{fila}")
        valores = [("D", r["magos"], "0"), ("E", r["venta"], "#,##0"), ("G", r["presupuesto"], "#,##0"),
                   ("I", r["cumplimiento"], "0.0%"), ("J", r["sss"], "+0.0%;-0.0%"), ("K", fz, "0.0%"),
                   ("L", int((de_casa["cumplimiento"] >= 1).sum()), "0")]
        for col, val, fmt in valores:
            color = TEXTO
            if col == "I":
                color = VERDE if (val or 0) >= 1 else ROJO
            elif col == "J":
                color = VERDE if (_num(val) or 0) >= 0 else ROJO
            elif col == "K":
                color = VERDE if (val or 0) > BONUS_FIDELIZACION["umbral"] else TEXTO
            merge = {"E": f"E{fila}:F{fila}", "G": f"G{fila}:H{fila}", "L": f"L{fila}:M{fila}"}.get(col)
            _celda(ws, f"{col}{fila}", _num(val), color=color, fill=fondo, fmt=fmt, h="center", merge=merge,
                   bold=col == "I")
        ws.row_dimensions[fila].height = 20
    n = len(casas)
    # Los gráficos leen la misma tabla: casa (B), % ppto (I) y SSS (J)
    cats = Reference(ws, min_col=2, min_row=26, max_row=25 + n)
    orden = list(casas["casa"])
    g1 = _grafico_casas(ws, "% de presupuesto por casa", cats, Reference(ws, min_col=9, min_row=26, max_row=25 + n),
                        orden, "0%")
    g2 = _grafico_casas(ws, "SSS por casa", cats, Reference(ws, min_col=10, min_row=26, max_row=25 + n), orden,
                        "+0.0%;-0.0%")
    ws.add_chart(g1, "B33")
    ws.add_chart(g2, "H33")

    # Mejores magos y los que necesitan magia
    base = 50
    _celda(ws, f"B{base}", "✦ TOP 10 MAGOS", size=15, bold=True, color=ORO_CLARO, font=SERIF, merge=f"B{base}:G{base}")
    _celda(ws, f"H{base}", "✦ NECESITAN MÁS MAGIA", size=15, bold=True, color=ORO_CLARO, font=SERIF,
           merge=f"H{base}:M{base}")
    top = rank.head(10)
    bajos = rank.sort_values("cumplimiento").head(10)
    for c0, tabla, medallas in (("B", top, True), ("H", bajos, False)):
        i0 = ord(c0) - 64
        cols = [get_column_letter(i0 + j) for j in range(6)]
        for col, t in zip(cols, ["#", "MAGO", "", "CASA", "PTS", "% PPTO"]):
            _celda(ws, f"{col}{base + 2}", t, size=9, bold=True, color=ORO, fill=BURDEOS, h="center")
        ws.merge_cells(f"{cols[1]}{base + 2}:{cols[2]}{base + 2}")
        for k, (_, r) in enumerate(tabla.iterrows()):
            fila = base + 3 + k
            fondo = PANEL if k % 2 == 0 else PANEL_2
            puesto = int(rank.index[rank["tienda"] == r["tienda"]][0]) + 1
            medalla = {1: "🥇", 2: "🥈", 3: "🥉"}.get(puesto) if medallas else None
            _celda(ws, f"{cols[0]}{fila}", medalla or puesto, color=ORO_CLARO, fill=fondo, h="center", bold=True)
            _celda(ws, f"{cols[1]}{fila}", r["tienda"], fill=fondo, bold=True, merge=f"{cols[1]}{fila}:{cols[2]}{fila}")
            _celda(ws, f"{cols[3]}{fila}", r["casa"], color=CLARO[r["casa"]],
                   fill=fondo, h="center", bold=True)
            _celda(ws, f"{cols[4]}{fila}", int(r["puntos"]), color=ORO_CLARO, fill=fondo, h="center", bold=True)
            _celda(ws, f"{cols[5]}{fila}", _num(r["cumplimiento"]), fmt="0.0%", fill=fondo, h="center", bold=True,
                   color=VERDE if (r["cumplimiento"] or 0) >= 1 else ROJO)
            ws.row_dimensions[fila].height = 19

    # Competencias
    base = 65
    _celda(ws, f"B{base}", "✦ COMPETENCIAS", size=15, bold=True, color=ORO_CLARO, font=SERIF, merge=f"B{base}:M{base}")
    for col, t, m in (("B", "COMPETENCIA", "B:E"), ("F", "FECHAS", "F:G"), ("H", "ESTADO", "H:I"),
                      ("J", "CUMPLIMIENTO RETAIL", "J:K"), ("L", "SE PAGA", "L:M")):
        a, b = m.split(":")
        _celda(ws, f"{col}{base + 2}", t, size=9, bold=True, color=ORO, fill=BURDEOS, h="center",
               merge=f"{a}{base + 2}:{b}{base + 2}")
    for k, c in enumerate(competencias):
        fila = base + 3 + k
        fondo = PANEL if k % 2 == 0 else PANEL_2
        ok = c["cumplimiento"] is not None and c["cumplimiento"] >= PAGA_DESDE
        _celda(ws, f"B{fila}", c["nombre"], bold=True, fill=fondo, merge=f"B{fila}:E{fila}", font=SERIF)
        _celda(ws, f"F{fila}", f"{c['inicio']:%d/%m} – {c['fin']:%d/%m}", fill=fondo, h="center", merge=f"F{fila}:G{fila}")
        _celda(ws, f"H{fila}", c["estado"], fill=fondo, h="center", merge=f"H{fila}:I{fila}",
               color=ORO_CLARO if c["estado"] == "En curso" else TEXTO)
        _celda(ws, f"J{fila}", _num(c["cumplimiento"]), fmt="0.0%", fill=fondo, h="center", bold=True,
               merge=f"J{fila}:K{fila}", color=(VERDE if ok else ROJO) if c["cumplimiento"] is not None else GRIS)
        _celda(ws, f"L{fila}", _num(c["factor"]) if c["cumplimiento"] is not None else None, fmt="0%", fill=fondo,
               h="center", merge=f"L{fila}:M{fila}", color=VERDE if ok else GRIS, bold=True)
        ws.row_dimensions[fila].height = 20
    _celda(ws, f"B{base + 9}", f"Se paga el premio desde {PAGA_DESDE:.0%} de cumplimiento del canal Retail. "
                               "Desempate: mejor % de cumplimiento.", size=9, color=GRIS, merge=f"B{base + 9}:M{base + 9}")
    _celda(ws, f"B{base + 11}", "✨ Travesura realizada ✨", size=12, color=ORO, font=SERIF, h="center",
           merge=f"B{base + 11}:M{base + 11}")
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:N78"


# ---------- Hojas de detalle ----------
COLUMNAS = [  # (campo, título, ancho, formato)
    ("puesto", "#", 6, "0"), ("tienda", "MAGO (TIENDA)", 28, None), ("casa", "CASA", 15, None),
    ("puntos", "PUNTOS", 10, "0"), ("pts_venta", "PTS VENTA", 10, "0"), ("pts_bonus", "PTS BONUS", 10, "0"),
    ("tramos_ganados", "TRAMOS GANADOS", 10, "0"), ("venta", "VENTA S/", 14, "#,##0"),
    ("presupuesto", "PRESUPUESTO S/", 15, "#,##0"), ("cumplimiento", "% PPTO", 10, "0.0%"),
    ("venta_ly", "VENTA AÑO PASADO S/", 15, "#,##0"), ("sss", "SSS", 10, "+0.0%;-0.0%"),
    ("tickets", "TICKETS", 10, "#,##0"), ("tickets_fidelizados", "TICKETS FIDELIZ.", 11, "#,##0"),
    ("fidelizacion", "% FIDELIZ.", 10, "0.0%"), ("ticket_prom", "TICKET PROM. S/", 11, "#,##0.0"),
    ("upt", "UPT", 8, "0.00"),
]
MEDALLAS = {1: "D4AF37", 2: "C0C0C0", 3: "CD7F32"}


def _detalle(ws, r: pd.DataFrame, titulo: str, subtitulo: str, casa: str | None):
    ws.sheet_view.showGridLines = False
    color = _hex(casa) if casa else BURDEOS
    ws.sheet_properties.tabColor = color
    ultima = get_column_letter(len(COLUMNAS))
    for i, (_, _, ancho, _) in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    # Cabecera: escudo de la casa (o el cáliz) sobre la banda de color
    for f in (1, 2, 3):
        ws.row_dimensions[f].height = 24
    _pinta(ws, f"A1:{ultima}3", color)
    _pinta(ws, f"A4:{ultima}4", ORO)
    ws.row_dimensions[4].height = 4
    if casa:
        ws.add_image(_escudo(casa, 84), "A1")
    else:
        ws.add_image(_imagen(ASSETS / "excel-banda.jpg", 440, "JPEG"), "A1")
    x0 = "C" if casa else "F"
    _celda(ws, f"{x0}1", titulo, size=20, bold=True, color="FFFFFF" if casa else ORO_CLARO, font=SERIF,
           merge=f"{x0}1:{ultima}2", v="bottom")
    _celda(ws, f"{x0}3", subtitulo, size=10, color="F5ECD0", merge=f"{x0}3:{ultima}3", v="top")

    # Tabla
    cab = 6
    for i, (_, t, _, _) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=cab, column=i, value=t)
        c.font = Font(name="Calibri", size=9, bold=True, color=ORO_CLARO)
        c.fill = _fill(NOCHE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color=ORO))
    ws.row_dimensions[cab].height = 32
    t = r.reset_index(drop=True).assign(puesto=lambda x: range(1, len(x) + 1))
    linea = Side(style="thin", color="E6DCC3")
    for k, (_, x) in enumerate(t.iterrows()):
        fila = cab + 1 + k
        fondo = "FFFFFF" if k % 2 == 0 else PERGAMINO
        for i, (campo, _, _, fmt) in enumerate(COLUMNAS, start=1):
            v = x[campo]
            v = _num(v) if fmt else v
            c = ws.cell(row=fila, column=i, value=int(v) if fmt == "0" and v is not None else v)
            c.fill = _fill(fondo)
            c.border = Border(bottom=linea)
            c.font = Font(name="Calibri", size=10, color="1E1E1E")
            c.alignment = Alignment(horizontal="left" if campo in ("tienda",) else "center", vertical="center")
            if fmt:
                c.number_format = fmt
        ws.cell(row=fila, column=2).font = Font(name="Calibri", size=10, bold=True, color="1E1E1E")
        ws.cell(row=fila, column=3).font = Font(name="Calibri", size=10, bold=True,
                                                color=_hex(x["casa"]) if x["casa"] != "Ownership" else "A47E10")
        ws.cell(row=fila, column=4).font = Font(name="Calibri", size=11, bold=True, color=BURDEOS)
        p = int(x["puesto"])
        if p in MEDALLAS:
            ws.cell(row=fila, column=1).fill = _fill(MEDALLAS[p])
            ws.cell(row=fila, column=1).font = Font(name="Calibri", size=10, bold=True, color="1E1E1E")
        ws.row_dimensions[fila].height = 18
    fin = cab + len(t)
    if len(t):
        col = {campo: get_column_letter(i) for i, (campo, *_) in enumerate(COLUMNAS, start=1)}
        verde, rojo = Font(color="1E8A4C", bold=True), Font(color="C0392B", bold=True)
        rng = lambda c: f"{col[c]}{cab + 1}:{col[c]}{fin}"
        ws.conditional_formatting.add(rng("cumplimiento"), CellIsRule(operator="greaterThanOrEqual", formula=["1"], font=verde))
        ws.conditional_formatting.add(rng("cumplimiento"), CellIsRule(operator="lessThan", formula=["1"], font=rojo))
        ws.conditional_formatting.add(rng("sss"), CellIsRule(operator="greaterThanOrEqual", formula=["0"], font=verde))
        ws.conditional_formatting.add(rng("sss"), CellIsRule(operator="lessThan", formula=["0"], font=rojo))
        ws.conditional_formatting.add(rng("fidelizacion"), CellIsRule(
            operator="greaterThan", formula=[str(BONUS_FIDELIZACION["umbral"])], font=verde))
        ws.conditional_formatting.add(rng("puntos"), DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                 color=ORO, showValue=True))
        ws.auto_filter.ref = f"A{cab}:{ultima}{fin}"
    _celda(ws, f"A{fin + 2}", "✨ Travesura realizada ✨", size=11, color=ORO, font=SERIF, h="center",
           merge=f"A{fin + 2}:{ultima}{fin + 2}")
    ws.freeze_panes = f"C{cab + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{cab}:{cab}"


def excel_concurso(rank: pd.DataFrame, dias: pd.DataFrame, competencias: list[dict], desde: date, hasta: date,
                   simulacion: bool = False, consultado: datetime | None = None) -> bytes:
    """rank: ranking de magos de las fechas elegidas · dias: venta diaria de esas fechas."""
    periodo = f"del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"
    extra = (" · simulación con la venta real de set–oct" if simulacion else "") + (
        f" · actualizado {consultado:%d/%m %H:%M}" if consultado else "")
    wb = Workbook()
    wb.active.title = "Dashboard"  # antes de los gráficos: sus referencias llevan el nombre de la hoja
    _dashboard(wb.active, rank, dias, competencias, f"Resumen {periodo}{extra}")
    _detalle(wb.create_sheet("Acumulado"), rank, "Copa del Mejor Mago · Acumulado",
             f"Los {len(rank)} magos {periodo}{extra}", None)
    for casa in CASAS:
        r = rank[rank["casa"] == casa]
        res = resumen_casas(r).iloc[0] if len(r) else None
        sub = (f"{CASAS[casa]['animal']} · {len(r)} magos · {res['cumplimiento']:.1%} de presupuesto · "
               f"SSS {res['sss']:+.1%} · {periodo}" if res is not None and pd.notna(res["sss"]) else
               f"{len(r)} magos · {periodo}")
        _detalle(wb.create_sheet(casa), r, f"Casa {casa}", sub, casa)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
